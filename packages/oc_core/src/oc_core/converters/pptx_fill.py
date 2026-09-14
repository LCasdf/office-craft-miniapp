"""Build a simple pptx from outline pages (tpl_basic_01 — programmatic master)."""

from __future__ import annotations

from pathlib import Path
from typing import Any


class PptxError(Exception):
    def __init__(self, message: str, *, kind: str = "invalid"):
        super().__init__(message)
        self.kind = kind


def render_outline_pptx(pages: list[dict[str, Any]], dest: Path) -> Path:
    """Title slide + bullet slides. Ceiling: text-only 16:9; no charts/images."""
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except ImportError as e:
        raise PptxError("python-pptx not installed", kind="dependency") from e

    if not pages:
        raise PptxError("empty pages", kind="invalid")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Cover
    cover = pages[0]
    slide = prs.slides.add_slide(blank)
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(1.4))
    p = title_box.text_frame.paragraphs[0]
    p.text = str(cover.get("title") or "封面")[:80]
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    p.alignment = PP_ALIGN.CENTER
    bullets = cover.get("bullets") or []
    if bullets:
        sub = slide.shapes.add_textbox(Inches(0.8), Inches(3.8), Inches(11.7), Inches(1.0))
        sp = sub.text_frame.paragraphs[0]
        sp.text = str(bullets[0])[:120]
        sp.font.size = Pt(20)
        sp.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
        sp.alignment = PP_ALIGN.CENTER

    for page in pages[1:]:
        s = prs.slides.add_slide(blank)
        hb = s.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(12), Inches(0.9))
        hp = hb.text_frame.paragraphs[0]
        hp.text = str(page.get("title") or "内容")[:80]
        hp.font.size = Pt(28)
        hp.font.bold = True
        hp.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

        body = s.shapes.add_textbox(Inches(0.9), Inches(1.6), Inches(11.5), Inches(5.2))
        tf = body.text_frame
        tf.word_wrap = True
        bl = page.get("bullets") or ["（空）"]
        for i, line in enumerate(bl):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = f"• {str(line)[:120]}"
            para.font.size = Pt(18)
            para.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
            para.space_after = Pt(10)

    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(dest))
    return dest
