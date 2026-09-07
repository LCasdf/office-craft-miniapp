#!/usr/bin/env python3
"""PPT template POC: build a 1-slide master + fill placeholders with python-pptx.

Run (no permanent dep):
  uv run --with python-pptx python scripts/ppt_poc/fill_demo.py

Writes scripts/ppt_poc/out/demo.pptx and prints Go/No-Go checklist hints.
"""

from __future__ import annotations

from pathlib import Path

OUT = Path(__file__).resolve().parent / "out"
OUT.mkdir(parents=True, exist_ok=True)


def main() -> int:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)

    # Title placeholder (text box with {{title}})
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.2))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "{{title}}"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    p.alignment = PP_ALIGN.CENTER

    sub = slide.shapes.add_textbox(Inches(0.8), Inches(3.8), Inches(11.7), Inches(0.8))
    sp = sub.text_frame.paragraphs[0]
    sp.text = "{{subtitle}}"
    sp.font.size = Pt(20)
    sp.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
    sp.alignment = PP_ALIGN.CENTER

    template_path = OUT / "template_blank.pptx"
    prs.save(template_path)

    # Fill: replace tokens in a copy
    filled = Presentation(str(template_path))
    mapping = {"{{title}}": "Office Craft POC", "{{subtitle}}": "大纲二编 → 模板填充可行"}
    for s in filled.slides:
        for shape in s.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    for k, v in mapping.items():
                        if k in run.text:
                            run.text = run.text.replace(k, v)
    demo = OUT / "demo.pptx"
    filled.save(demo)
    print(f"template: {template_path}")
    print(f"filled:   {demo}")
    print("POC smoke: OK (python-pptx can create + fill text tokens)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
